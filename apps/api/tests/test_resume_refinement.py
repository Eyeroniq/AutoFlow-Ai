"""Multi-agent resume refinement: each agent's output schema and extra checks, the PDF layout
measurements, the guardrails the assembler applies, the generated PDF/DOCX, and the API flow
(upload -> refine -> stages -> download -> email to the owner only) with a scripted LLM.
The real models are exercised by the verification runs, not here."""

import io
import json
import re

import pymupdf
import pytest
from docx import Document
from flowforge_engine.jsonschema_lite import validate

from app.services import resume_agents as ra
from app.services import resume_refinement as rr
from app.services.resume_documents import build_docx, build_pdf
from app.services.resume_layout import analyze_pdf, layout_issues
from tests import resume_samples
from tests.test_files_api import upload

PARSED = {
    "contact_info": {"name": "Priya Nair", "email": "priya.nair@example.com", "phone": "+91 98765 43210",
                     "location": "Pune, India", "links": ["github.com/pnair"]},
    "summary": "Backend engineer with 6 years of experience.",
    "experience": [
        {"company": "FinServe", "title": "Senior Engineer", "dates": "Mar 2021 - Present", "bullets": [
            "Responsible for the backend of the onboarding platform.",
            "Led the migration of a Django monolith to 7 microservices, cutting deploy time from 3 hours to 20 minutes.",
            "Helped with various production incidents.",
        ]},
        {"company": "ShopKart", "title": "Software Engineer", "dates": "Jul 2018 - Feb 2021", "bullets": [
            "Developed REST APIs for the checkout service.",
            "Wrote tests.",
        ]},
    ],
    "education": [{"institution": "SPPU", "degree": "B.E. Computer Engineering", "dates": "2014 - 2018", "details": ""}],
    "skills": ["Python", "FastAPI", "PostgreSQL", "Redis"],
}


def content_for(parsed=PARSED, flag=("e0.b0", "e0.b2", "e1.b1")):
    bullets = []
    for bid, info in ra.bullet_index(parsed).items():
        if bid in flag:
            bullets.append({"id": bid, "flags": ["weak_opening_verb", "no_metric"], "explanation": f"Weak start: {info['text'][:20]}",
                            "rewrite": {"e0.b0": "Own the backend of the onboarding platform serving [N] merchants.",
                                        "e0.b2": "Resolved various production incidents on call [N incidents].",
                                        "e1.b1": "Wrote unit tests reaching [X%] coverage."}[bid]})
        else:
            bullets.append({"id": bid, "flags": [], "explanation": "", "rewrite": ""})
    return {"overall": "Mixed.", "summary_review": "Generic.", "bullets": bullets}


def rewrite_for(parsed=PARSED, content=None, reorder=False):
    content = content or content_for(parsed)
    new = {b["id"]: (b["rewrite"] or ra.bullet_index(parsed)[b["id"]]["text"]) for b in content["bullets"]}
    jobs = []
    for i, job in enumerate(parsed["experience"]):
        bullets = [{"text": new[f"e{i}.b{j}"], "source_id": f"e{i}.b{j}"} for j in range(len(job["bullets"]))]
        jobs.append({**{k: job[k] for k in ("company", "title", "dates")}, "bullets": bullets[::-1] if reorder else bullets})
    return {"summary": "Backend engineer who ships reliable Python services.", "experience": jobs, "skills": parsed["skills"], "notes": "Applied feedback."}


def parsed_pdf() -> bytes:
    """A PDF containing PARSED's text (the scripted Parser Agent's answer can then be checked against it)."""
    return build_pdf(PARSED)


ATS_OK = {"score": 92, "summary": "Clean.", "issues": []}
MATCH_OK = {
    "match_score": 70, "summary": "Decent fit; Kubernetes is missing.",
    "matched_keywords": [{"keyword": "Python", "evidence": "FinServe"}, {"keyword": "Redis", "evidence": "skills"}],
    "missing_keywords": [{"keyword": "Kubernetes", "importance": "high", "suggestion": "True gap."}],
    "requirement_matches": [{"requirement": "Operate services in production", "strength": "strong", "bullet_ids": ["e0.b2"], "note": "On call."}],
}


# --- Output schemas ----------------------------------------------------------------------------


def test_the_parsed_resume_schema_accepts_the_structure_and_rejects_a_flattened_one():
    assert validate(PARSED, ra.PARSED_SCHEMA) == []
    broken = json.loads(json.dumps(PARSED))
    broken["experience"][0]["bullets"] = "one long string"
    del broken["skills"]
    problems = validate(broken, ra.PARSED_SCHEMA)
    assert any("skills" in p for p in problems) and any("bullets" in p for p in problems)


@pytest.mark.parametrize(("schema", "good", "bad"), [
    (ra.ATS_SCHEMA, {"score": 40, "summary": "Bad.", "issues": [
        {"category": "multi_column", "severity": "blocker", "title": "Two columns", "explanation": "x" * 30, "fix": "Use one column", "evidence": "p1"}]},
     {"score": 140, "summary": "", "issues": [{"category": "colour", "severity": "fatal", "title": "t", "explanation": "e", "fix": "f"}]}),
    (ra.CONTENT_SCHEMA, content_for(), {"overall": "x", "summary_review": "", "bullets": [{"id": "e0.b0", "flags": ["boring"], "explanation": "", "rewrite": ""}]}),
    (ra.JOB_MATCH_SCHEMA, MATCH_OK, {"match_score": "high", "summary": "", "matched_keywords": [{"keyword": "x"}], "missing_keywords": [], "requirement_matches": []}),
    (ra.REWRITE_SCHEMA, rewrite_for(), {"summary": "s", "experience": [{"company": "c", "title": "t", "dates": "d", "bullets": ["plain string"]}], "skills": [], "notes": ""}),
])
def test_each_agents_schema_accepts_its_shape_and_rejects_garbage(schema, good, bad):
    assert validate(good, schema) == []
    assert validate(bad, schema) != []


# --- Extra checks ------------------------------------------------------------------------------


def test_parser_checks_catch_missing_names_empty_jobs_and_invented_bullets():
    assert ra.check_parsed(PARSED) == []
    nameless = json.loads(json.dumps(PARSED))
    nameless["contact_info"]["name"] = ""
    nameless["experience"][1]["bullets"] = []
    problems = ra.check_parsed(nameless)
    assert any("name is empty" in p for p in problems) and any("no bullets" in p for p in problems)
    source = "Priya Nair\n" + " ".join(b for j in PARSED["experience"] for b in j["bullets"])
    assert ra.check_parsed(PARSED, source) == []
    invented = json.loads(json.dumps(PARSED))
    invented["experience"][0]["bullets"][0] = "Single-handedly rebuilt the entire payments platform from scratch."
    assert any("don't appear in the resume text" in p for p in ra.check_parsed(invented, source))


def test_content_checks_require_every_bullet_valid_ids_and_a_rewrite_for_each_flag():
    bullets = ra.bullet_index(PARSED)
    assert ra.check_content(content_for(), bullets) == []
    data = content_for()
    data["bullets"] = data["bullets"][:-1]
    assert any("not reviewed" in p for p in ra.check_content(data, bullets))
    data = content_for()
    data["bullets"][0]["id"] = "e9.b9"
    assert any("isn't one of the given bullet ids" in p for p in ra.check_content(data, bullets))
    data = content_for()
    data["bullets"][0]["rewrite"] = ""
    assert any("has no rewrite" in p for p in ra.check_content(data, bullets))
    data = content_for()
    data["bullets"][1]["rewrite"] = "Led it better."
    assert any("no flags" in p for p in ra.check_content(data, bullets))


def test_invented_numbers_and_wrong_tense_are_soft_problems():
    bullets = ra.bullet_index(PARSED)
    data = content_for()
    data["bullets"][0]["rewrite"] = "Own the backend serving 5 million merchants."  # 5 million: invented
    data["bullets"][4]["rewrite"] = "Write tests reaching [X%] coverage."  # e1.b1 is a PAST role: "Write"
    data["bullets"][4]["flags"] = ["weak_opening_verb"]
    problems = ra.check_content(data, bullets)
    soft = [p for p in problems if p.startswith(ra.SOFT)]
    assert len(soft) == 2 and len(soft) == len(problems)
    assert any("invents the number" in p for p in soft) and any("PAST role" in p for p in soft)


def test_number_and_tense_helpers():
    assert ra.numbers_in("Cut costs by 40% and $1,200 in 3 weeks; saw [X%] gains") == {"40", "1200", "3"}
    assert ra.invented_numbers("Led 7 engineers, saving [N] hours", "Led 7 engineers") == set()
    assert ra.invented_numbers("Led 9 engineers", "Led 7 engineers") == {"9"}
    assert ra.wrong_tense("Develop APIs", current=False) and not ra.wrong_tense("Developed APIs", current=False)
    assert not ra.wrong_tense("Led a team", current=False) and not ra.wrong_tense("Develop APIs", current=True)
    assert ra.is_current({"dates": "Mar 2021 - Present"}) and not ra.is_current({"dates": "2018 - 2020"})


def test_padding_a_rewrite_with_invented_clauses_is_caught_but_a_stronger_verb_is_not():
    original = "Collaborated across teams to make great products"
    assert ra.padded("Collaborated across teams to deliver [N] products, ensuring cohesive design and functionality.", original)
    assert ra.padded("Managed customer service and cash transactions for [N] daily patrons.", "Served customers and handled cash.")
    assert not ra.padded("Analyzed sales data utilizing Excel and SQL across [N] datasets.", "Used Excel and some SQL to look at sales data.")
    assert not ra.padded("Architected and maintained the Python and FastAPI backend for the merchant onboarding platform",
                         "Responsible for the backend of the merchant onboarding platform, built with Python and FastAPI.")
    data = content_for()
    data["bullets"][0]["rewrite"] = "Own the backend of the onboarding platform, ensuring seamless merchant growth and delighted stakeholders."
    assert any("adds claims" in p for p in ra.check_content(data, ra.bullet_index(PARSED)))


def test_unknown_ids_are_soft_and_sanitized_and_a_good_early_reply_beats_a_worse_retry():
    bullets = ra.bullet_index(PARSED)
    data = content_for()
    data["bullets"].append({"id": "e0.b9", "flags": [], "explanation": "", "rewrite": ""})
    problems = ra.check_content(data, bullets)
    assert problems and all(p.startswith(ra.SOFT) for p in problems)  # accepted after retries, then cleaned
    assert [b["id"] for b in ra.sanitize_content(data, bullets)["bullets"]] == list(bullets)
    match = json.loads(json.dumps(MATCH_OK))
    match["requirement_matches"][0]["bullet_ids"] = ["e0.b2", "e7.b7"]
    assert ra.sanitize_job_match(match, bullets)["requirement_matches"][0]["bullet_ids"] == ["e0.b2"]


class SequenceLLM:
    is_mock = True
    name = "mock"

    def __init__(self, *replies: str):
        self.replies = list(replies)

    async def generate(self, system_prompt, user_prompt, model, temperature, max_tokens):
        return self.replies.pop(0)


async def test_run_agent_falls_back_to_the_earlier_reply_that_had_only_soft_problems():
    from flowforge_engine import ExecutionServices, NodeContext

    bullets = ra.bullet_index(PARSED)
    soft = content_for()
    soft["bullets"][4]["rewrite"] = "Write unit tests reaching [X%] coverage."  # e1.b1 is a past role: a soft tense problem
    hard = content_for()
    hard["bullets"] = hard["bullets"][:-2]  # a later reply that forgets bullets: a hard problem
    services = ExecutionServices()
    services._llm["mock"] = SequenceLLM(json.dumps(soft), json.dumps(hard), json.dumps(hard))
    context = NodeContext(workflow_id="t", execution_id="t", services=services)
    run = await ra.run_agent(context, "content", "review", ra.CONTENT_SCHEMA, provider="mock", fallback=[],
                             check=lambda d: ra.check_content(d, bullets))
    assert run.error is None and run.output == soft and len(run.attempts) == 3
    assert any("PAST role" in w for w in run.warnings)
    services._llm["mock"] = SequenceLLM(*[json.dumps(hard)] * 3)
    failed = await ra.run_agent(context, "content", "review", ra.CONTENT_SCHEMA, provider="mock", fallback=[],
                                check=lambda d: ra.check_content(d, bullets))
    assert failed.output is None and "failed validation after 3 attempts" in failed.error


def test_job_match_checks_cite_real_bullets():
    bullets = ra.bullet_index(PARSED)
    assert ra.check_job_match(MATCH_OK, bullets) == []
    bad = json.loads(json.dumps(MATCH_OK))
    bad["requirement_matches"][0]["bullet_ids"] = ["e7.b7"]
    assert any("unknown bullet ids" in p for p in ra.check_job_match(bad, bullets))
    bad["requirement_matches"][0]["bullet_ids"] = []
    bad["requirement_matches"][0]["note"] = ""
    assert any("no bullet ids" in p for p in ra.check_job_match(bad, bullets))
    bad["requirement_matches"][0]["note"] = "Shown by the B.E. in Computer Engineering (education)."  # met by education, not a bullet
    assert ra.check_job_match(bad, bullets) == []


def test_rewrite_checks_keep_the_structure_every_bullet_once_and_no_new_skills():
    assert ra.check_rewrite(rewrite_for(), PARSED, content_for()) == []
    fewer = rewrite_for()
    fewer["experience"].pop()
    assert any("must have 2 entries" in p for p in ra.check_rewrite(fewer, PARSED))
    dropped = rewrite_for()
    dropped["experience"][0]["bullets"].pop()
    assert any("must appear exactly once" in p for p in ra.check_rewrite(dropped, PARSED))
    moved = rewrite_for()
    moved["experience"][0]["bullets"][0]["source_id"] = "e1.b0"
    assert any("belongs to another job" in p for p in ra.check_rewrite(moved, PARSED))
    skills = rewrite_for()
    skills["skills"] = [*PARSED["skills"], "Kubernetes"]
    assert any("Kubernetes" in p for p in ra.check_rewrite(skills, PARSED))


# --- Layout ------------------------------------------------------------------------------------------


def test_layout_analysis_finds_what_makes_a_resume_hostile_to_an_ats(tmp_path):
    clean = tmp_path / "clean.pdf"
    clean.write_bytes(resume_samples.priya())
    hostile = tmp_path / "hostile.pdf"
    hostile.write_bytes(resume_samples.marcus())
    assert layout_issues(analyze_pdf(str(clean))) == []
    layout = analyze_pdf(str(hostile))
    found = {i["category"]: i["severity"] for i in layout_issues(layout)}
    assert found == {"multi_column": "blocker", "table": "blocker", "image_or_icon": "warning", "header_footer": "blocker"}
    assert layout["pages"][0]["columns"] == 2 and layout["contact_in_header_footer"]
    assert "MY JOURNEY" in layout["nonstandard_headers"]


def test_a_measured_blocker_the_model_missed_still_reaches_the_report():
    measured = [{"category": "multi_column", "severity": "blocker", "title": "Two-column layout", "explanation": "e" * 30, "fix": "one column"}]
    merged = rr.merge_ats_issues({"score": 80, "summary": "ok", "issues": [
        {"category": "length", "severity": "warning", "title": "Long", "explanation": "x" * 30, "fix": "cut"}]}, measured)
    assert [(i["category"], i["source"]) for i in merged["issues"]] == [("multi_column", "layout"), ("length", "agent")]  # blockers first
    again = rr.merge_ats_issues({"score": 80, "summary": "ok", "issues": [
        {"category": "multi_column", "severity": "blocker", "title": "Columns", "explanation": "x" * 30, "fix": "f"}]}, measured)
    assert len(again["issues"]) == 1  # the model already reported it


# --- Assembling and the guardrails -----------------------------------------------------------------------


def test_assembly_applies_the_rewrite_and_reports_before_and_after_per_bullet():
    result = rr.assemble(PARSED, ATS_OK, content_for(), None, rewrite_for())
    by_id = {c["source_id"]: c for c in result["changes"]}
    assert by_id["e0.b0"]["changed"] and by_id["e0.b0"]["before"].startswith("Responsible for")
    assert by_id["e0.b0"]["after"].startswith("Own the backend") and by_id["e0.b0"]["flags"] == ["weak_opening_verb", "no_metric"]
    assert not by_id["e0.b1"]["changed"]
    stats = result["stats"]
    assert (stats["bullets_total"], stats["bullets_rewritten"], stats["bullets_flagged"], stats["jd"]) == (5, 3, 3, False)
    assert result["resume"]["contact_info"] == PARSED["contact_info"] and result["resume"]["education"] == PARSED["education"]
    assert result["resume"]["experience"][0]["bullets"][0].startswith("Own the backend")


def test_facts_the_agent_changes_are_restored_from_the_parsed_resume():
    sneaky = rewrite_for()
    sneaky["experience"][0].update(company="Google", title="CTO", dates="1999 - 2030")
    job = rr.assemble(PARSED, ATS_OK, content_for(), None, sneaky)["resume"]["experience"][0]
    assert (job["company"], job["title"], job["dates"]) == ("FinServe", "Senior Engineer", "Mar 2021 - Present")


def test_guardrails_remove_invented_numbers_fix_tense_and_keep_untouched_bullets_verbatim():
    rewrite = rewrite_for()
    bullets = rewrite["experience"][0]["bullets"]
    bullets[0]["text"] = "Own the backend serving 5 million merchants."  # invented number -> the Content Agent's suggestion
    bullets[1]["text"] = "Led the migration of a Django monolith to 7 microservices; deploys now take minutes."  # untouched bullet reworded
    rewrite["experience"][1]["bullets"][0]["text"] = "Develop REST APIs for checkout."  # past role, present tense
    result = rr.assemble(PARSED, ATS_OK, content_for(), None, rewrite)
    by_id = {c["source_id"]: c for c in result["changes"]}
    assert by_id["e0.b0"]["after"] == "Own the backend of the onboarding platform serving [N] merchants."
    assert by_id["e0.b1"]["after"] == by_id["e0.b1"]["before"]
    assert by_id["e1.b0"]["after"] == "Developed REST APIs for the checkout service."  # no suggestion: the original
    reasons = {c["source_id"]: c["reason"] for c in result["corrections"]}
    assert "invented number" in reasons["e0.b0"] and "stays as written" in reasons["e0.b1"] and "past-tense" in reasons["e1.b0"]
    assert result["stats"]["corrections"] == 3


def test_with_a_job_description_untouched_bullets_may_be_tailored_and_bullets_may_move():
    rewrite = rewrite_for(reorder=True)
    target = next(b for b in rewrite["experience"][0]["bullets"] if b["source_id"] == "e0.b1")
    target["text"] = "Led the migration of a Django monolith to 7 microservices on AWS, cutting deploy time from 3 hours to 20 minutes."
    result = rr.assemble(PARSED, ATS_OK, content_for(), MATCH_OK, rewrite)
    by_id = {c["source_id"]: c for c in result["changes"]}
    assert "on AWS" in by_id["e0.b1"]["after"] and result["stats"]["bullets_reordered"] >= 2
    assert result["stats"]["keywords_matched"] == 2 and result["stats"]["keywords_total"] == 3 and result["stats"]["missing_high"] == ["Kubernetes"]
    assert by_id["e0.b2"]["job_requirements"] == ["Operate services in production"]


def test_without_a_job_description_the_original_bullet_order_is_kept():
    result = rr.assemble(PARSED, ATS_OK, content_for(), None, rewrite_for(reorder=True))
    assert [c["source_id"] for c in result["changes"] if c["job"] == 0] == ["e0.b0", "e0.b1", "e0.b2"]
    assert result["stats"]["bullets_reordered"] == 0


def test_the_email_summary_counts_what_changed():
    row = rr.ResumeRefinement(version=2, status="success", filename="cv.pdf", result_json=rr.assemble(
        PARSED, {"score": 40, "summary": "x", "issues": [
            {"category": "multi_column", "severity": "blocker", "title": "Columns", "explanation": "e" * 30, "fix": "f"},
            {"category": "table", "severity": "blocker", "title": "Table", "explanation": "e" * 30, "fix": "f"},
            {"category": "length", "severity": "warning", "title": "Too long", "explanation": "e" * 30, "fix": "f"}]},
        content_for(), MATCH_OK, rewrite_for()))
    text = rr.email_summary(row)
    assert text.startswith("Rewrote 3 of 5 bullets, fixed 2 ATS formatting issues")
    assert "matched 2 of 3 job-description keywords" in text and "Still for you to look at: Too long" in text
    assert "3 placeholders in square brackets" in text
    assert rr.document_name(row, "docx") == "cv-refined-v2.docx"


# --- The generated documents ---------------------------------------------------------------------------------


def final_resume():
    return rr.assemble(PARSED, ATS_OK, content_for(), None, rewrite_for())["resume"]


def test_the_refined_pdf_is_text_single_column_and_passes_the_ats_checks_itself(tmp_path):
    path = tmp_path / "refined.pdf"
    path.write_bytes(build_pdf(final_resume()))
    layout = analyze_pdf(str(path))
    assert layout_issues(layout) == [] and layout["pages"][0]["columns"] == 1
    with pymupdf.open(path) as pdf:
        text = pdf[0].get_text()
    assert "Priya Nair" in text and "EXPERIENCE" in text and "Own the backend of the onboarding platform" in text
    assert "priya.nair@example.com" in text and "FinServe" in text and "Python, FastAPI" in text


def test_the_refined_docx_has_real_headings_bullets_and_no_tables():
    document = Document(io.BytesIO(build_docx(final_resume())))
    texts = [p.text for p in document.paragraphs]
    assert "PRIYA NAIR" not in texts and texts[0] == "Priya Nair"
    assert {"SUMMARY", "EXPERIENCE", "EDUCATION", "SKILLS"} <= set(texts)
    assert sum(1 for p in document.paragraphs if p.style.name == "List Bullet") == 5
    assert document.tables == [] and not document.inline_shapes
    assert any("Own the backend of the onboarding platform" in t for t in texts)


# --- The API flow, with a scripted LLM -------------------------------------------------------------------------


class ScriptedLLM:
    """Answers each agent by the role named in its prompt, with valid JSON built from the prompt's own contents."""

    is_mock = True
    name = "mock"

    def __init__(self, fail_on: str | None = None):
        self.calls: list[str] = []
        self.fail_on = fail_on

    async def generate(self, system_prompt, user_prompt, model, temperature, max_tokens):
        role = re.search(r"You are the ([A-Za-z&/ -]+?)(?: Agent| in a resume)", user_prompt).group(1)
        self.calls.append(role)
        if self.fail_on and self.fail_on in role:
            return "I'm sorry, I cannot do that."
        if role.startswith("Parser"):
            return json.dumps(PARSED)
        if role.startswith("ATS"):
            has_columns = "2 text column(s)" in user_prompt
            return json.dumps({"score": 40 if has_columns else 95, "summary": "Layout reviewed.", "issues": [] if not has_columns else [
                {"category": "multi_column", "severity": "blocker", "title": "Two columns", "explanation": "Columns scramble the reading order for an ATS.",
                 "fix": "Use one column.", "evidence": "measured"}]})
        if role.startswith("Content"):
            return json.dumps(content_for())
        if role.startswith("Job-Match"):
            return json.dumps(MATCH_OK)
        return json.dumps(rewrite_for())


@pytest.fixture
def scripted(monkeypatch):
    llm = ScriptedLLM()
    original = rr.build_execution_services

    async def build(db, user, **kwargs):
        services = await original(db, user, **kwargs)
        services._llm["mock"] = llm
        return services

    monkeypatch.setattr(rr, "build_execution_services", build)
    return llm


@pytest.fixture
def inline(refine_queue, session_factory):
    async def run(refinement_id):
        await rr.run_refinement(session_factory, refinement_id)

    refine_queue.inline = run
    return run


async def start(client, user, pdf, jd="", name="cv.pdf"):
    uploaded = await upload(client, user, pdf, filename=name)
    assert uploaded.status_code == 201, uploaded.text
    response = await client.post("/api/resume-refinements", json={"file_id": uploaded.json()["id"], "job_description": jd}, headers=user.headers)
    assert response.status_code == 202, response.text
    detail = await client.get(f"/api/resume-refinements/{response.json()['id']}", headers=user.headers)
    return detail.json()


async def test_a_refinement_runs_every_agent_and_stores_each_stage_with_its_raw_replies(client, user, files_dir, scripted, inline):
    body = await start(client, user, parsed_pdf(), jd=resume_samples.PRIYA_JD)
    assert body["status"] == "success" and body["version"] == 1 and body["has_job_description"]
    stages = {s["key"]: s for s in body["stages"]}
    assert list(stages) == ["extract", "parse", "ats", "content", "job_match", "rewrite"]
    assert all(s["status"] == "success" for s in stages.values())
    for key in ("parse", "ats", "content", "job_match", "rewrite"):
        stage = stages[key]
        assert stage["attempts"] and stage["attempts"][0]["raw"].startswith("{") and stage["output"] and stage["prompt"]
        assert stage["provider_used"] == "mock" and stage["duration_ms"] is not None
    assert stages["extract"]["output"]["layout"]["page_count"] == 1
    assert sorted(set(scripted.calls)) == ["ATS Compatibility", "Content & Impact", "Job-Match", "Parser", "Rewrite"]
    assert body["result"]["stats"]["bullets_rewritten"] == 3 and body["result"]["resume"]["contact_info"]["name"] == "Priya Nair"


async def test_the_job_match_agent_only_runs_with_a_job_description(client, user, files_dir, scripted, inline):
    body = await start(client, user, parsed_pdf())
    stages = {s["key"]: s for s in body["stages"]}
    assert stages["job_match"]["status"] == "skipped" and "No job description" in stages["job_match"]["note"]
    assert "Job-Match" not in scripted.calls and body["result"]["stats"]["jd"] is False


async def test_the_ats_agent_sees_the_measured_layout_and_measured_blockers_are_reported(client, user, files_dir, scripted, inline, monkeypatch):
    monkeypatch.setattr(ra, "check_parsed", lambda data, source="": [])  # the scripted parse isn't this PDF's text
    body = await start(client, user, resume_samples.marcus(), name="marcus.pdf")
    ats = next(s for s in body["stages"] if s["key"] == "ats")["output"]
    categories = {i["category"]: i["source"] for i in ats["issues"]}
    assert categories["multi_column"] == "agent" and {"table", "header_footer", "image_or_icon"} <= set(categories)
    assert categories["table"] == "layout" and ats["issues"][0]["severity"] == "blocker"
    assert body["result"]["stats"]["ats_blockers"] == 3 and body["result"]["stats"]["ats_fixed_by_reformat"] == 4


async def test_refining_the_same_file_again_makes_the_next_version(client, user, files_dir, scripted, inline):
    pdf = parsed_pdf()
    first, second = await start(client, user, pdf), await start(client, user, pdf)
    assert (first["version"], second["version"]) == (1, 2)
    listing = (await client.get("/api/resume-refinements", headers=user.headers)).json()
    assert [r["version"] for r in listing] == [2, 1]


async def test_an_agent_that_cannot_produce_valid_json_fails_the_run_with_its_stage_visible(client, user, files_dir, monkeypatch, inline):
    llm = ScriptedLLM(fail_on="Content")
    original = rr.build_execution_services

    async def build(db, user_, **kwargs):
        services = await original(db, user_, **kwargs)
        services._llm["mock"] = llm
        return services

    monkeypatch.setattr(rr, "build_execution_services", build)
    body = await start(client, user, parsed_pdf())
    assert body["status"] == "failed" and "Content & Impact Agent failed" in body["error_message"]
    stages = {s["key"]: s for s in body["stages"]}
    assert stages["parse"]["status"] == "success" and stages["content"]["status"] == "failed"
    assert len(stages["content"]["attempts"]) == 3 and stages["content"]["attempts"][0]["raw"].startswith("I'm sorry")
    assert stages["rewrite"]["status"] == "pending" and body["result"] is None


async def test_a_scanned_pdf_with_no_text_is_refused_with_advice(client, user, files_dir, scripted, inline):
    blank = pymupdf.open()
    blank.new_page()
    body = await start(client, user, blank.tobytes())
    assert body["status"] == "failed" and "no selectable text" in body["error_message"]
    assert next(s for s in body["stages"] if s["key"] == "parse")["status"] == "pending"


async def test_only_pdfs_are_accepted_and_only_your_own_files(client, user, user_factory, files_dir, refine_queue):
    text = await upload(client, user, b"just text", filename="cv.txt", content_type="text/plain")
    refused = await client.post("/api/resume-refinements", json={"file_id": text.json()["id"]}, headers=user.headers)
    assert refused.status_code == 415
    pdf = await upload(client, user, parsed_pdf())
    other = await user_factory()
    stranger = await client.post("/api/resume-refinements", json={"file_id": pdf.json()["id"]}, headers=other.headers)
    assert stranger.status_code == 404


async def test_a_queue_outage_fails_the_request_cleanly(client, user, files_dir, refine_queue):
    refine_queue.fail = True
    pdf = await upload(client, user, parsed_pdf())
    response = await client.post("/api/resume-refinements", json={"file_id": pdf.json()["id"]}, headers=user.headers)
    assert response.status_code == 503
    listing = (await client.get("/api/resume-refinements", headers=user.headers)).json()
    assert listing[0]["status"] == "failed" and "Could not queue" in listing[0]["error_message"]


async def test_download_gives_a_real_pdf_and_docx_and_waits_for_the_run(client, user, files_dir, scripted, refine_queue, inline):
    refine_queue.inline = None  # queued, not run yet
    pdf = await upload(client, user, parsed_pdf(), filename="Priya CV.pdf")
    started = (await client.post("/api/resume-refinements", json={"file_id": pdf.json()["id"]}, headers=user.headers)).json()
    pending = await client.get(f"/api/resume-refinements/{started['id']}/download", headers=user.headers)
    assert pending.status_code == 409
    await inline(refine_queue.queued[0])
    got_pdf = await client.get(f"/api/resume-refinements/{started['id']}/download?format=pdf", headers=user.headers)
    got_docx = await client.get(f"/api/resume-refinements/{started['id']}/download?format=docx", headers=user.headers)
    assert got_pdf.status_code == got_docx.status_code == 200
    assert got_pdf.headers["content-disposition"] == 'attachment; filename="Priya-CV-refined-v1.pdf"'
    with pymupdf.open(stream=got_pdf.content, filetype="pdf") as doc:
        assert "Own the backend" in doc[0].get_text()
    assert Document(io.BytesIO(got_docx.content)).paragraphs[0].text == "Priya Nair"


async def test_emailing_sends_both_files_to_the_owner_only_and_records_it(client, user, files_dir, scripted, inline, monkeypatch):
    sent = []

    class Mail:
        async def send_email(self, email):
            sent.append(email)
            return {"message_id": "<abc@test>", "status": "sent"}

    from flowforge_engine import ExecutionServices

    monkeypatch.setattr(ExecutionServices, "email", lambda self, name: Mail())
    body = await start(client, user, parsed_pdf(), jd=resume_samples.PRIYA_JD)
    first = await client.post(f"/api/resume-refinements/{body['id']}/email", headers=user.headers)
    assert first.status_code == 201, first.text
    (mail,) = sent
    assert mail.to == [user.email] and mail.cc == [] and mail.bcc == []  # the account's own address, nobody else
    assert [a.filename for a in mail.attachments] == ["cv-refined-v1.pdf", "cv-refined-v1.docx"]
    assert mail.attachments[0].data()[:4] == b"%PDF" and mail.attachments[1].data()[:2] == b"PK"
    assert "Rewrote 3 of 5 bullets" in mail.body and "matched 2 of 3 job-description keywords" in mail.body
    record = first.json()
    assert record["to_email"] == user.email and record["version"] == 1 and record["message_id"] if "message_id" in record else True
    await client.post(f"/api/resume-refinements/{body['id']}/email", headers=user.headers)
    detail = (await client.get(f"/api/resume-refinements/{body['id']}", headers=user.headers)).json()
    assert len(detail["emails"]) == 2 and detail["emails_sent"] == 2
    # There is no way to name a recipient: extra fields are ignored and the address stays the owner's.
    await client.post(f"/api/resume-refinements/{body['id']}/email", json={"to": "someone@else.com"}, headers=user.headers)
    assert sent[-1].to == [user.email]


async def test_emailing_needs_a_finished_run_and_ownership(client, user, user_factory, files_dir, scripted, refine_queue, inline):
    refine_queue.inline = None
    pdf = await upload(client, user, parsed_pdf())
    started = (await client.post("/api/resume-refinements", json={"file_id": pdf.json()["id"]}, headers=user.headers)).json()
    assert (await client.post(f"/api/resume-refinements/{started['id']}/email", headers=user.headers)).status_code == 409
    await inline(refine_queue.queued[0])
    stranger = await user_factory()
    assert (await client.post(f"/api/resume-refinements/{started['id']}/email", headers=stranger.headers)).status_code == 404
    assert (await client.get(f"/api/resume-refinements/{started['id']}", headers=stranger.headers)).status_code == 404
    assert (await client.delete(f"/api/resume-refinements/{started['id']}", headers=user.headers)).status_code == 204


# --- The agents as pipeline nodes -------------------------------------------------------------------------------


async def test_the_resume_nodes_chain_into_the_same_refinement(db_session, user, files_dir):
    from flowforge_engine import GraphNode, NodeContext, execute_node

    from app.models.user import User
    from app.services.credentials import build_execution_services

    owner = await db_session.get(User, user.id)
    services = await build_execution_services(db_session, owner)
    llm = ScriptedLLM()
    services._llm["mock"] = llm
    context = NodeContext(workflow_id="w", execution_id="e", services=services)
    text = " ".join(b for j in PARSED["experience"] for b in j["bullets"]) + " Priya Nair backend engineer " * 3

    async def run(node_type, **config):
        result = await execute_node(GraphNode(id=node_type, type=node_type, config=config), context)
        assert result.status.value == "success", result.error
        return result.output

    parsed = (await run("resume_parse", text=text))["data"]
    ats = (await run("resume_ats", parsed=parsed))["data"]
    content = (await run("resume_content", parsed=parsed))["data"]
    skipped = await run("resume_match", parsed=parsed, job_description="")
    assert skipped["matched"] is False and skipped["data"] is None and "Job-Match" not in llm.calls
    match = (await run("resume_match", parsed=parsed, job_description=resume_samples.PRIYA_JD))["data"]
    final = await run("resume_rewrite", parsed=parsed, ats=ats, content=content, match=match, job_description=resume_samples.PRIYA_JD)
    assert final["stats"]["bullets_rewritten"] == 3 and final["resume"]["contact_info"]["name"] == "Priya Nair"
    assert "Priya Nair" in final["text"] and "EXPERIENCE" in final["text"]
    bad = await execute_node(GraphNode(id="r", type="resume_content", config={"parsed": "not a resume"}), context)
    assert bad.status.value != "success" and "Resume Parse" in bad.error
